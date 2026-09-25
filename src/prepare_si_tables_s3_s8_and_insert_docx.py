#!/usr/bin/env python3
from __future__ import annotations

import json
import re
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document


REPO = Path(__file__).resolve().parents[1]
SI_TABLE_DIR = REPO / "reports" / "main_manuscript_assets" / "SI" / "tables"
SI_DOCX = REPO / "reports" / "main_manuscript_assets" / "SI" / "Supporting Information.docx"
SI_DOCX_OUT = REPO / "reports" / "main_manuscript_assets" / "SI" / "Supporting Information_with_tables_S3_S8.docx"


def parse_feature_name(name: str) -> dict:
	lag = re.search(r"_lag(\d+)$", name)
	roll = re.search(r"_roll(\d+)$", name)
	if lag:
		base = name[: lag.start()]
		window = int(lag.group(1))
		return {
			"feature_name": name,
			"base_variable": base,
			"transformation": "lag",
			"parameter": window,
			"formula": f"{base}(t-{window})",
		}
	if roll:
		base = name[: roll.start()]
		window = int(roll.group(1))
		return {
			"feature_name": name,
			"base_variable": base,
			"transformation": "rolling_mean",
			"parameter": window,
			"formula": f"mean({base}, window={window})",
		}
	return {
		"feature_name": name,
		"base_variable": name,
		"transformation": "raw",
		"parameter": "",
		"formula": f"{name}(t)",
	}


def build_s3_feature_dictionary() -> pd.DataFrame:
	npz_map = {
		"H1": REPO / "features" / "ulsan_H1_features.npz",
		"H3": REPO / "features" / "ulsan_H3_features.npz",
		"H5": REPO / "features" / "ulsan_H5_features.npz",
		"H5_v2": REPO / "features" / "ulsan_H5_features_v2.npz",
	}
	feat_by_h = {}
	all_features: set[str] = set()
	for h, path in npz_map.items():
		arr = np.load(path, allow_pickle=True)["feature_names"]
		feats = [str(x) for x in arr.tolist()]
		feat_by_h[h] = set(feats)
		all_features.update(feats)

	rows = []
	for feat in sorted(all_features):
		row = parse_feature_name(feat)
		row.update(
			{
				"in_H1": int(feat in feat_by_h["H1"]),
				"in_H3": int(feat in feat_by_h["H3"]),
				"in_H5": int(feat in feat_by_h["H5"]),
				"in_H5_v2": int(feat in feat_by_h["H5_v2"]),
				"horizon_dependent": int(
					len({feat in feat_by_h[k] for k in feat_by_h}) > 1
				),
			}
		)
		rows.append(row)
	return pd.DataFrame(rows)


def build_s4_preprocessing_rules() -> pd.DataFrame:
	rows = [
		{
			"model_family": "Persistence",
			"missing_handling": "No model fitting; uses lagged observed TNout directly",
			"scaling": "None",
			"categorical_handling": "N/A",
			"outlier_clipping": "None",
			"notes": "Deterministic baseline for each horizon",
		},
		{
			"model_family": "ElasticNet",
			"missing_handling": "Uses prepared NPZ features from preprocessing pipeline",
			"scaling": "StandardScaler in sklearn Pipeline",
			"categorical_handling": "N/A",
			"outlier_clipping": "None in model code",
			"notes": "Configured in run_linear_alarm.py and save_fold_predictions.py",
		},
		{
			"model_family": "HGBR",
			"missing_handling": "Uses prepared NPZ features; tree model handles split-based robustness",
			"scaling": "None",
			"categorical_handling": "N/A",
			"outlier_clipping": "None in training script",
			"notes": "Hyperparameters tuned with Optuna",
		},
		{
			"model_family": "TCN (BCR-TCN)",
			"missing_handling": "Uses prepared NPZ tensors",
			"scaling": "Handled in model input preparation",
			"categorical_handling": "N/A",
			"outlier_clipping": "None explicitly documented",
			"notes": "Trained for alarm-focused ranking metrics",
		},
		{
			"model_family": "Hybrid rank ensemble",
			"missing_handling": "Inherits from base model predictions",
			"scaling": "Rank-based combination (not feature scaling)",
			"categorical_handling": "N/A",
			"outlier_clipping": "N/A",
			"notes": "Guarded policy with budget cap and tau-specific ranking",
		},
	]
	return pd.DataFrame(rows)


def build_s5_model_config_summary() -> pd.DataFrame:
	run_cfg_path = REPO / "results" / "hybrid_rank_v2_runs" / "20260220_131820_H5_hybrid_rank_v2" / "run_config.json"
	run_cfg = json.loads(run_cfg_path.read_text())
	rows = [
		{
			"model": "Persistence",
			"input_window_length": "lag=H",
			"feature_set": "TNout lag baseline",
			"objective_or_loss": "baseline deterministic",
			"early_stopping": "N/A",
			"output_type": "point prediction",
		},
		{
			"model": "ElasticNet",
			"input_window_length": "tabular lag/rolling features",
			"feature_set": "NPZ engineered feature set",
			"objective_or_loss": "L1/L2-regularized regression",
			"early_stopping": "No",
			"output_type": "point prediction",
		},
		{
			"model": "HGBR",
			"input_window_length": "tabular lag/rolling features",
			"feature_set": "NPZ engineered feature set",
			"objective_or_loss": "least-squares boosting",
			"early_stopping": "No explicit early stop; fixed max_iter",
			"output_type": "point prediction",
		},
		{
			"model": "TCN (BCR-TCN)",
			"input_window_length": "sequence window (per training config)",
			"feature_set": "temporal feature tensors",
			"objective_or_loss": "budget-constrained ranking/classification objective",
			"early_stopping": "Configured in training routine",
			"output_type": "risk score / alarm ranking",
		},
		{
			"model": "Hybrid rank v2",
			"input_window_length": "inherits from base models",
			"feature_set": "base model scores",
			"objective_or_loss": run_cfg.get("rank_objective", "rank objective"),
			"early_stopping": "N/A",
			"output_type": "alarm ranking",
		},
	]
	return pd.DataFrame(rows)


def build_s6_hyperparameters_final() -> pd.DataFrame:
	rows: list[dict] = []

	rows.append(
		{
			"model": "ElasticNet",
			"horizon": "H1/H3/H5",
			"alpha": 0.1,
			"l1_ratio": 0.5,
			"max_iter": 20000,
			"random_state": 42,
			"source": "src/save_fold_predictions.py",
		}
	)

	hgbr_files = [
		REPO / "results" / "hgbr" / "hgbr_optuna_H3.json",
		REPO / "results" / "hgbr" / "hgbr_optuna_H5.json",
		REPO / "results" / "hgbr" / "hgbr_optuna_H5_v2.json",
	]
	for path in hgbr_files:
		obj = json.loads(path.read_text())
		h = obj.get("h", "")
		params = obj.get("best_params", {})
		row = {"model": "HGBR", "horizon": f"H{h}", "source": str(path.relative_to(REPO))}
		row.update(params)
		rows.append(row)

	for path in [
		REPO / "results" / "metrics" / "bcr_tcn_v11_H5_v2_meta.json",
		REPO / "results" / "metrics" / "bcr_tcn_H5_v2_meta.json",
	]:
		obj = json.loads(path.read_text())
		row = {
			"model": path.stem.replace("_meta", ""),
			"horizon": f"H{obj.get('h', '')}",
			"source": str(path.relative_to(REPO)),
		}
		cfg = obj.get("cfg", {})
		for k, v in cfg.items():
			if isinstance(v, (int, float, str, bool)):
				row[k] = v
		rows.append(row)

	return pd.DataFrame(rows)


def build_s7_exceedance_prevalence() -> pd.DataFrame:
	path = REPO / "results" / "paper_artifacts" / "20260220_133021_ACS_Uncertainty_Decision" / "tables" / "Table2_exceedance_prevalence_random_baseline.csv"
	df = pd.read_csv(path)
	return df.rename(
		columns={
			"horizon": "H",
			"tau": "tau_mgL",
			"n_test": "N",
			"events": "E",
			"prevalence": "pi",
			"random_recall_at_r005": "random_recall_r005",
			"random_recall_at_r010": "random_recall_r010",
		}
	)


def build_s8_fold_rows() -> pd.DataFrame:
	path = REPO / "results" / "paper_artifacts" / "20260220_133021_ACS_Uncertainty_Decision" / "tables" / "Table5_key_operating_points_fold_rows.csv"
	df = pd.read_csv(path)
	model_map = {
		"persistence": "Persistence",
		"enet": "ElasticNet",
		"hgbr": "HGBR",
		"tcn": "TCN",
		"hybrid_rank_v2": "Hybrid (guarded)",
		"hybrid_paper_fixed": "Hybrid (paper-fixed)",
	}
	df["model"] = df["model"].map(lambda x: model_map.get(x, x))
	return df.sort_values(["horizon", "tau", "budget_r", "model", "fold"]).reset_index(drop=True)


def insert_df_as_table_after_paragraph(doc: Document, paragraph, df: pd.DataFrame) -> None:
	table = doc.add_table(rows=df.shape[0] + 1, cols=df.shape[1])
	try:
		table.style = "Table Grid"
	except KeyError:
		pass

	for col_idx, col_name in enumerate(df.columns):
		table.rows[0].cells[col_idx].text = str(col_name)

	for row_idx in range(df.shape[0]):
		for col_idx in range(df.shape[1]):
			val = df.iat[row_idx, col_idx]
			if isinstance(val, float):
				text = f"{val:.6g}"
			else:
				text = str(val)
			table.rows[row_idx + 1].cells[col_idx].text = text

	paragraph._p.addnext(table._tbl)


def insert_tables_into_docx(doc_path: Path, out_path: Path, table_files: dict[str, Path]) -> None:
	doc = Document(str(doc_path))
	placeholders = {f"S{i}": f"[Table S{i} placed here]" for i in range(3, 9)}

	for p in list(doc.paragraphs):
		txt = p.text.strip()
		for table_id, marker in placeholders.items():
			if marker in txt:
				df = pd.read_csv(table_files[table_id])
				insert_df_as_table_after_paragraph(doc, p, df)
				break

	doc.save(str(out_path))


def main() -> None:
	SI_TABLE_DIR.mkdir(parents=True, exist_ok=True)

	outputs: dict[str, Path] = {
		"S3": SI_TABLE_DIR / "TableS3_feature_dictionary_formulas.csv",
		"S4": SI_TABLE_DIR / "TableS4_preprocessing_rules_by_model_family.csv",
		"S5": SI_TABLE_DIR / "TableS5_model_configuration_summary.csv",
		"S6": SI_TABLE_DIR / "TableS6_hyperparameters_final.csv",
		"S7": SI_TABLE_DIR / "TableS7_exceedance_prevalence_by_horizon_threshold.csv",
		"S8": SI_TABLE_DIR / "TableS8_fold_level_precision_recall_selected_operating_points.csv",
	}

	build_s3_feature_dictionary().to_csv(outputs["S3"], index=False)
	build_s4_preprocessing_rules().to_csv(outputs["S4"], index=False)
	build_s5_model_config_summary().to_csv(outputs["S5"], index=False)
	build_s6_hyperparameters_final().to_csv(outputs["S6"], index=False)
	build_s7_exceedance_prevalence().to_csv(outputs["S7"], index=False)
	build_s8_fold_rows().to_csv(outputs["S8"], index=False)

	insert_tables_into_docx(SI_DOCX, SI_DOCX_OUT, outputs)

	print("Generated tables:")
	for key, value in outputs.items():
		print(f"- {key}: {value.relative_to(REPO)}")
	print(f"Updated SI docx: {SI_DOCX_OUT.relative_to(REPO)}")


if __name__ == "__main__":
	main()
